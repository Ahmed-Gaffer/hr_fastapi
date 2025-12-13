import tkinter as tk
from tkinter import filedialog
import os
from docx import Document
import re

TARGET_EXTENSIONS = [".py", ".rkt", ".txt", ".json", ".css"]

def clean_text(text):
    text = text.replace("\x00", "")
    text = re.sub(r"[\x01-\x08\x0B-\x0C\x0E-\x1F]", "", text)
    return text

def read_file_content(filepath):
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    return None

def main():
    root = tk.Tk()
    root.withdraw()

    all_files = []
    while True:
        files = filedialog.askopenfilenames(
            title="اختر ملفات (اضغط Cancel للانتهاء)",
            filetypes=[("Text/Code files", "*.py *.rkt *.txt *.json *.css"), ("All files", "*.*")]
        )
        if not files:  # لو ضغط Cancel
            break
        all_files.extend(files)

    if not all_files:
        print("❌ لم يتم اختيار أي ملف")
        return

    doc = Document()
    doc.add_heading("تجميع محتويات الملفات", level=0)

    for path in all_files:
        if not path.endswith(tuple(TARGET_EXTENSIONS)):
            continue
        content = read_file_content(path)
        if content:
            content = clean_text(content)
            doc.add_heading(f"ملف: {path}", level=1)
            doc.add_paragraph(content)
            doc.add_page_break()
        else:
            doc.add_heading(f"⚠️ لم أستطع قراءة الملف: {path}", level=2)

    output_path = os.path.join(os.getcwd(), "merged_files.docx")
    doc.save(output_path)
    print(f"✅ تم حفظ كل الملفات بمحتواها في: {output_path}")

if __name__ == "__main__":
    main()