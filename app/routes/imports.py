from fastapi import APIRouter
from fastapi.responses import JSONResponse
from app.services.import_excel import import_employees_from_excel
import re
import uuid
from io import BytesIO
from datetime import datetime, date
from typing import Dict, Any, Optional

import openpyxl
from fastapi import APIRouter, UploadFile, File, Depends, HTTPException, Query
from sqlmodel import Session, select

from app.database import get_session
from app.models.employee import Employee
from app.models.site import Site
from app.models.cost_center import CostCenter
from app.models.salary_record import SalaryRecord, PayrollCategory
from app.models.salary_component import SalaryComponent, ComponentType
from app.services.professional_salary_calculator import ProfessionalSalaryCalculator
from app.services.smart_importer import SmartSalaryImporter

router = APIRouter(prefix="/import", tags=["import"])

HEADER_NORMALIZE_MAP = {
    "الكود القديم": "legacy_code",
    "الاسم": "name",
    "موقع العمل": "site",
    "الادارة / المشروع": "cost_center",
    "الوظيفة": "role",
    "تاريخ التأمين": "hire_date",
    "اساسي": "basic_salary",
    "اعاشة": "allowance_food",
    "سفر": "allowance_travel",
    "مواصلات": "allowance_transport",
    "بدلات اخري": "allowance_position",
    "ساعات عمل نهارية": "extra_hours",
    "إجمالي قيمة ساعات العمل الإضافية": "extra_hours_value",
    "مكافأة أداء": "bonus_performance",
    "مكافأة مشروع": "bonus_project",
    "استحقاقات أخرى": "other_earnings",
    "تأمين اجتماعي (العامل)": "deduction_insurance_social",
    "تأمين صحي": "deduction_insurance_health",
    "غياب": "deduction_absence",
    "جزاءات": "deduction_penalties",
    "خصومات أخرى": "deduction_other",
    "أقساط": "loan_deduction",
    "سلف راتب": "advance_salary",
    "صافي الراتب المستحق": "net_salary",
    "مرتبات شهر": "salary_date",
    "عدد أيام الحضور": "attendance_days",
    "عدد ايام الغياب": "absence_days",
    "الرقم القومي": "national_id",
    "الرقم التأميني": "insurance_number",
    "فئة": "payroll_category",
}

def norm_header(h: Optional[str]) -> Optional[str]:
    if not h:
        return None
    return HEADER_NORMALIZE_MAP.get(str(h).strip(), str(h).strip())

def parse_num(v) -> float:
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    cleaned = re.sub(r"[^\d\.\-]", "", s)
    if not cleaned or cleaned == ".":
        return 0.0
    try:
        return float(cleaned)
    except:
        return 0.0

def parse_date(v) -> Optional[date]:
    if not v:
        return None
    if isinstance(v, datetime):
        return v.date()
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except:
            pass
    return None

class BulkImporter:
    def __init__(self, file_bytes: bytes, session: Session):
        self.wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        self.session = session
        self.report = {
            "sites": {"imported": 0, "errors": []},
            "cost_centers": {"imported": 0, "errors": []},
            "employees": {"created": 0, "updated": 0, "errors": []},
            "salaries": {"imported": 0, "errors": []},
            "warnings": [],
            "mapping_legacy_to_new": {},
            "status": "pending"
        }
        self._gen_counter = 0
        self._seen = set()

    def _gen_code(self) -> str:
        self._gen_counter += 1
        ts = datetime.utcnow().strftime("%y%m%d%H%M%S")
        return f"EMP{ts}{self._gen_counter:03d}"

    def _find_or_create_site(self, name: Optional[str]):
        if not name:
            return None
        name = str(name).strip()
        s = self.session.exec(select(Site).where(Site.name == name)).first()
        if s:
            return s
        s = Site(name=name)
        self.session.add(s)
        self.session.flush()
        self.report["sites"]["imported"] += 1
        return s

    def _find_or_create_cost_center(self, name: Optional[str], site_obj: Optional[Site]):
        if not name:
            return None
        name = str(name).strip()
        cc = self.session.exec(select(CostCenter).where(CostCenter.name == name)).first()
        if cc:
            return cc
        cc = CostCenter(code=str(uuid.uuid4())[:8], name=name, site_id=(site_obj.id if site_obj else None))
        self.session.add(cc)
        self.session.flush()
        self.report["cost_centers"]["imported"] += 1
        return cc

    def _find_employee(self, legacy: Optional[str], nat_id: Optional[str], ins_num: Optional[str]) -> Optional[Employee]:
        for val in [legacy, nat_id, ins_num]:
            if val:
                q = self.session.exec(select(Employee).where(Employee.code == str(val).strip())).first()
                if q:
                    return q
        return None

    def import_file(self, commit: bool = False) -> Dict[str, Any]:
        try:
            if "الرواتب" not in self.wb.sheetnames:
                self.report["salaries"]["errors"].append("شيت 'الرواتب' غير موجود")
                self.report["status"] = "failed"
                return self.report

            ws = self.wb["الرواتب"]
            headers = {}
            for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)), 1):
                key = norm_header(cell)
                if key:
                    headers[key] = idx - 1

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                rowd = {k: row[idx] if idx < len(row) else None for k, idx in headers.items()}

                name = rowd.get("name")
                if not name:
                    self.report["salaries"]["errors"].append(f"صف {row_idx}: الاسم مفقود")
                    continue

                salary_date_raw = rowd.get("salary_date")
                if not salary_date_raw:
                    self.report["salaries"]["errors"].append(f"صف {row_idx}: تاريخ المرتب مفقود")
                    continue

                sdate = parse_date(salary_date_raw)
                if not sdate:
                    self.report["salaries"]["errors"].append(f"صف {row_idx}: خطأ في تحويل التاريخ")
                    continue

                year, month = sdate.year, sdate.month
                legacy = rowd.get("legacy_code")
                identifier = str(legacy or name).strip()
                
                if (identifier, year, month) in self._seen:
                    self.report["salaries"]["errors"].append(f"صف {row_idx}: تكرار موظف ونفس الشهر")
                    continue
                self._seen.add((identifier, year, month))

                site_obj = self._find_or_create_site(rowd.get("site"))
                cc_obj = self._find_or_create_cost_center(rowd.get("cost_center"), site_obj)

                emp = self._find_employee(legacy, rowd.get("national_id"), rowd.get("insurance_number"))
                if not emp:
                    new_code = self._gen_code()
                    emp = Employee(
                        code=new_code,
                        name=str(name).strip(),
                        role=str(rowd.get("role") or "").strip(),
                        site_id=(site_obj.id if site_obj else None),
                        cost_center_id=(cc_obj.id if cc_obj else None),
                        hire_date=parse_date(rowd.get("hire_date")),
                        national_id=str(rowd.get("national_id")).strip() if rowd.get("national_id") else None,
                        insurance_number=str(rowd.get("insurance_number")).strip() if rowd.get("insurance_number") else None,
                        base_salary=parse_num(rowd.get("basic_salary")),
                        status="نشط"
                    )
                    try:
                        self.session.add(emp)
                        self.session.flush()
                        self.report["employees"]["created"] += 1
                        if legacy:
                            self.report["mapping_legacy_to_new"][str(legacy).strip()] = emp.code
                    except Exception as e:
                        self.report["employees"]["errors"].append(f"صف {row_idx}: {str(e)}")
                        continue
                else:
                    if site_obj:
                        emp.site_id = site_obj.id
                    if cc_obj:
                        emp.cost_center_id = cc_obj.id
                    try:
                        self.session.add(emp)
                        self.session.flush()
                        self.report["employees"]["updated"] += 1
                    except Exception as e:
                        self.report["employees"]["errors"].append(f"صف {row_idx}: {str(e)}")
                        continue

                exists = self.session.exec(
                    select(SalaryRecord).where(
                        SalaryRecord.employee_id == emp.id,
                        SalaryRecord.salary_year == year,
                        SalaryRecord.salary_month == month
                    )
                ).first()
                if exists:
                    self.report["salaries"]["errors"].append(f"صف {row_idx}: سجل راتب موجود")
                    continue

                cat_str = str(rowd.get("payroll_category") or "الفئة_الأولى").strip()
                payroll_cat = PayrollCategory.CATEGORY_A
                if "ثانية" in cat_str or "B" in cat_str:
                    payroll_cat = PayrollCategory.CATEGORY_B

                try:
                    # تحقق إذا كان آخر صف في الملف
                    total_rows = sum(1 for _ in ws.iter_rows(min_row=2))
                    is_last_month = (row_idx == total_rows + 1)
                    
                    smart_importer = SmartSalaryImporter(self.session)
                    
                    salary_input = {
                        "employee_id": emp.id,
                        "year": year,
                        "month": month,
                        "payroll_category": payroll_cat,
                        "basic_salary": parse_num(rowd.get("basic_salary")),
                        "allowance_food": parse_num(rowd.get("allowance_food")),
                        "allowance_transport": parse_num(rowd.get("allowance_transport")),
                        "allowance_travel": parse_num(rowd.get("allowance_travel")),
                        "allowance_position": parse_num(rowd.get("allowance_position")),
                        "extra_hours": parse_num(rowd.get("extra_hours")),
                        "extra_hours_value": parse_num(rowd.get("extra_hours_value")),
                        "bonus_performance": parse_num(rowd.get("bonus_performance")),
                        "bonus_project": parse_num(rowd.get("bonus_project")),
                        "other_earnings": parse_num(rowd.get("other_earnings")),
                        "deduction_insurance_social_employee": parse_num(rowd.get("deduction_insurance_social")),
                        "deduction_insurance_health": parse_num(rowd.get("deduction_insurance_health")),
                        "deduction_absence": parse_num(rowd.get("deduction_absence")),
                        "deduction_penalties": parse_num(rowd.get("deduction_penalties")),
                        "deduction_other": parse_num(rowd.get("deduction_other")),
                        "advance_salary": parse_num(rowd.get("advance_salary")),
                        "loan_deduction": parse_num(rowd.get("loan_deduction")),
                        "attendance_days": int(parse_num(rowd.get("attendance_days") or 30)),
                        "absence_days": int(parse_num(rowd.get("absence_days") or 0)),
                        "payment_date": sdate,
                        "net_salary_from_file": parse_num(rowd.get("net_salary")),  # من الملف
                        "notes": rowd.get("notes")
                    }
                    
                    import_result = smart_importer.import_salary_record(salary_input, is_last_month=is_last_month)
                    
                    if import_result.get("status") == "error":
                        self.report["salaries"]["errors"].append(f"صف {row_idx}: {import_result.get('message')}")
                        continue
                    
                    # إذا في تحقق للشهر الأخير
                    if import_result.get("verification"):
                        self.report["last_month_verification"] = import_result.get("verification")
                        if import_result.get("warnings"):
                            self.report["warnings"].extend(import_result.get("warnings"))
                    
                    self.report["salaries"]["imported"] += 1

                except Exception as e:
                    self.report["salaries"]["errors"].append(f"صف {row_idx}: {str(e)}")
                    continue

            if commit:
                self.session.commit()
                self.report["status"] = "success"
            else:
                self.session.rollback()
                self.report["status"] = "dry-run"
        except Exception as e:
            try:
                self.session.rollback()
            except:
                pass
            self.report["status"] = "failed"
            self.report["error"] = str(e)
        
        return self.report

@router.post("/bulk")
async def import_bulk(
    file: UploadFile = File(...),
    commit: bool = Query(False),
    session: Session = Depends(get_session)
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Excel فقط")
    
    contents = await file.read()
    importer = BulkImporter(contents, session)
    report = importer.import_file(commit=commit)
    
    if report.get("status") == "failed":
        raise HTTPException(status_code=500, detail=report.get("error"))
    
    return report

@router.get("/salary/{salary_record_id}/calculate")
async def calculate_salary(salary_record_id: int, session: Session = Depends(get_session)):
    try:
        calc = ProfessionalSalaryCalculator(session)
        return calc.calculate(salary_record_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/salary/{salary_record_id}/component")
async def add_component(
    salary_record_id: int,
    component_type: str,
    description: str,
    quantity: float,
    unit_rate: float,
    is_deduction: bool = False,
    session: Session = Depends(get_session)
):
    try:
        comp = SalaryComponent(
            salary_record_id=salary_record_id,
            component_type=component_type,
            description=description,
            quantity=quantity,
            unit_rate=unit_rate,
            amount=quantity * unit_rate,
            is_deduction=is_deduction
        )
        session.add(comp)
        session.commit()
        return {"id": comp.id, "amount": comp.amount}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
